#!/usr/bin/env python3
"""把 report.md 转成 Word 文档（python-docx）。

支持：标题、段落、粗体、行内代码、代码块、表格、无序/有序列表。
用法：
  python3 md_to_docx.py report.md -o report.docx
"""

import argparse
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def set_cn_font(run, name='宋体'):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_inline(paragraph, text):
    """处理 **粗体** 和 `行内代码`。"""
    tokens = re.split(r'(\*\*.*?\*\*|`[^`]*`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            set_cn_font(run)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        else:
            run = paragraph.add_run(token)
            set_cn_font(run)


def is_separator(row):
    return all(re.fullmatch(r':?-{2,}:?', c.strip()) for c in row)


def parse_table(lines, idx):
    rows = []
    while idx < len(lines) and lines[idx].strip().startswith('|'):
        row = [c.strip() for c in lines[idx].strip().strip('|').split('|')]
        if not is_separator(row):
            rows.append(row)
        idx += 1
    return rows, idx


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('input')
    parser.add_argument('-o', '--output', required=True)
    args = parser.parse_args()

    lines = open(args.input, encoding='utf-8').read().splitlines()
    doc = Document()

    # 默认正文样式
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            run = p.add_run('\n'.join(code))
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            i += 1
            continue

        image = re.fullmatch(r'!\[(.*?)\]\((.*?)\)', stripped)
        if image:
            alt, path = image.groups()
            if not os.path.isabs(path):
                path = os.path.join(
                    os.path.dirname(os.path.abspath(args.input)), path)
            if os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(path, width=Inches(6.0))
                cap = doc.add_paragraph(alt)
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cap.runs:
                    run.italic = True
                    run.font.size = Pt(9)
            else:
                doc.add_paragraph(f'[图片缺失: {path}]')
            i += 1
            continue

        link = re.fullmatch(r'\[(.*?)\]\((.*?)\)', stripped)
        if link:
            p = doc.add_paragraph()
            p.add_run(link.group(1) + '  ')
            p.add_run(link.group(2)).font.size = Pt(9)
            i += 1
            continue

        if stripped.startswith('#'):
            level = len(stripped) - len(stripped.lstrip('#'))
            text = stripped[level:].strip()
            if level == 1:
                p = doc.add_heading(text, level=0)
            else:
                p = doc.add_heading(text, level=min(level - 1, 4))
            for run in p.runs:
                set_cn_font(run, '黑体')
            i += 1
            continue

        if stripped.startswith('|') and i + 1 < len(lines) \
                and lines[i + 1].strip().startswith('|'):
            rows, i = parse_table(lines, i)
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=ncols)
                table.style = 'Table Grid'
                for r, row in enumerate(rows):
                    for c in range(ncols):
                        cell = table.cell(r, c)
                        cell.text = ''
                        p = cell.paragraphs[0]
                        text = row[c] if c < len(row) else ''
                        add_inline(p, text)
                        for run in p.runs:
                            run.font.size = Pt(9)
                            if r == 0:
                                run.bold = True
            continue

        if re.match(r'^\s*[-*]\s+', line):
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, re.sub(r'^\s*[-*]\s+', '', line))
            i += 1
            continue

        if re.match(r'^\s*\d+\.\s+', line):
            p = doc.add_paragraph(style='List Number')
            add_inline(p, re.sub(r'^\s*\d+\.\s+', '', line))
            i += 1
            continue

        if stripped in ('', '---'):
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(21)
        add_inline(p, line)
        i += 1

    doc.save(args.output)
    print('saved', args.output)


if __name__ == '__main__':
    main()
